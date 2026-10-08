# %% Initial setup
import numpy as np
import pandas as pd
import matplotlib.pyplot as plt
import os
import pywt
from scipy.integrate import cumulative_trapezoid
from scipy.signal import butter, filtfilt, detrend, find_peaks, welch
import tkinter as tk
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment, Border, Side
from matplotlib.backends.backend_pdf import PdfPages




def TurnDuration(G, fs, th):
    valid_indices_LB = G.index[G.abs() > th]
    i_start_LB=valid_indices_LB[0]
    i_end_LB=valid_indices_LB[-1]
    gyro = G.iloc[valid_indices_LB[0]:valid_indices_LB[-1]]
    
    Duration = (i_end_LB - i_start_LB) / fs;   
        
    return [Duration, i_start_LB, i_end_LB, gyro] 

def TurnAngle(G, fs):
    sample=G.shape[0]
    duration=np.arange(0, sample)/fs
    turn_angle=np.trapezoid(G,duration)
    return np.abs(turn_angle)


def StepDetection(acc, gyro, fs, orientation, idx_i):
    b, a = butter(4, [0.05/(0.5 * fs), 9/(0.5 * fs)], btype='band')
    acc_filt = filtfilt(b, a, acc)
    gyro_filt = filtfilt(b, a, gyro)

    widths = np.arange(1, fs//2)
    cwt_mat, _ = pywt.cwt(gyro_filt, widths, 'gaus1')
    gait_envelope = np.mean(np.abs(cwt_mat[10:30, :]), axis=0)
    
    gyro_seg1 = gyro_filt[:idx_i*2]
    gyro_seg2 = gyro_filt[idx_i*2:]
    
    if orientation == "Right":
        sig1 = gyro_seg1
        sig2 = -gyro_seg2
    else:
        sig1 = -gyro_seg1
        sig2 = gyro_seg2
        
    height1 = np.max(np.abs(sig1)) * 0.2  
    prom1 = np.std(sig1) * 0.1  
    height2 = np.max(np.abs(sig2)) * 0.2 
    prom2 = np.std(sig2) * 0.1 
    
    step_peaks1, _ = find_peaks(sig1, distance=fs*0.2, prominence=prom1, height=height1)
    step_peaks2, _ = find_peaks(sig2, distance=fs*0.2, prominence=prom2, height=height2)
    
    step_peaks2 = step_peaks2 + (idx_i*2)
    step_peaks = np.sort(np.concatenate((step_peaks1, step_peaks2)))   
    
    valid_IC = []
    valid_FC = []
    validated_peaks = []
    
    min_zero_time=0.1
    min_zero_samples = int(min_zero_time * fs)
    validated_peaks.append(step_peaks[0])
    for i in range(1, len(step_peaks)):
        prev_peak = step_peaks[i-1]
        curr_peak = step_peaks[i]

        segment = gyro_filt[prev_peak:curr_peak]
        zero_th = 0.4*np.max(np.abs(segment))
        near_zero = np.abs(segment) < zero_th

        count = 0
        found_reset = False

        for nz in near_zero:
            if nz:
                count += 1
                if count >= min_zero_samples:
                    found_reset = True
                    break
            else:
                count = 0

        if found_reset:
            validated_peaks.append(curr_peak)
    
    for p in validated_peaks:
        win = int(fs * 0.2)
        start, end = max(0, p - win), min(len(gyro_filt), p + win)
        
        gyro_segment = gyro_filt[start:end]
        acc_segment = acc_filt[start:end]

        if len(gyro_segment) == 0: continue

        local_fc = np.argmin(np.abs(gyro_segment[:len(gyro_segment)//2])) + start
        
        peak_swing_local = np.argmax(np.abs(gyro_segment))
        post_swing_acc = acc_segment[peak_swing_local:]
        
        if len(post_swing_acc) > 0:
            local_ic = np.argmax(np.abs(post_swing_acc)) + peak_swing_local + start
        else:
            local_ic = p 

        valid_FC.append(local_fc)
        valid_IC.append(local_ic)

    return np.array(valid_IC), np.array(valid_FC), gait_envelope, gyro_filt


# Clear console and close all the plots
os.system('cls' if os.name == 'nt' else 'clear')
plt.close('all')

dim_legend='5'

codes = set()
folder=Path("Dati_sensori_CSV")

# %% Call of the patients 

for file in folder.glob("*.csv"):
    name = file.name
    code = name.split("_")[0]
    codes.add(code)
    
conditions=["Condition2", "Condition4"]
tests=["Test5", "Test6"]
trials=["Trial1", "Trial2"]

pdf_filename = "Report2.pdf"
with PdfPages(pdf_filename) as pdf:
    for condition in conditions:
        for test in tests:
            for trial in trials:
                exist=0
                Table=pd.DataFrame()
                for code in codes:
                
                    filename_Feet = folder / f"{code}_TimeMeasure1_{condition}_{test}_{trial}_Feet.csv"
                    filename_LB   = folder / f"{code}_TimeMeasure1_{condition}_{test}_{trial}_LB.csv"
                
                    if filename_Feet.exists() & filename_LB.exists():
                        exist=exist+1    
                    else:
                        print(f"Not founded patient {code} \n")
                        continue
                    
                
                    Feet = pd.read_csv(filename_Feet)
                    LB = pd.read_csv(filename_LB)
                
                    
                    RF = Feet.iloc[:, :6]
                    LF = Feet.iloc[:, -6:]
                    
                    Freq_LB=100
                    Freq_F=200
                    
                    Len_LB=LB.shape[0]
                    Len_RF=RF.shape[0]
                    Len_LF=LF.shape[0]
                        
                    t_LB=np.arange(0,Len_LB)/Freq_LB
                    t_RF=np.arange(0,Len_RF)/Freq_F
                    t_LF=np.arange(0,Len_LF)/Freq_F
                    
                    # %% Lower back part
                    
                    g_y = LB.gyr_y;
                    th_LB = g_y.abs().max()*0.2
                    [duration_LB, i_start_LB, i_end_LB, gyro_moving] = TurnDuration(LB.gyr_y,Freq_LB, th_LB)
                    
                    MTV=gyro_moving.abs().mean() # Mean Turn Velocity (mean turn velocity)
                    PTV=gyro_moving.abs().max() # Peak Turn Velocity (max turn velocity)
            
                    angle_vect = cumulative_trapezoid(g_y, t_LB, initial=0);
                    idx_inversion = np.argmax(np.abs(angle_vect))
                    angle_LB=np.abs(angle_vect[idx_inversion])
                    
                    if angle_vect[idx_inversion]<0:
                        Orientation="Left"
                    else:
                        Orientation="Right"
                    
                
                    gyro_moving1 = gyro_moving[:idx_inversion-i_start_LB]; 
                    gyro_moving2=gyro_moving[idx_inversion-i_start_LB+1:]; 
                    MTV_ST1 = gyro_moving1.abs().mean();
                    MTV_ST2 = gyro_moving2.abs().mean();
                    PTV_ST1 = gyro_moving1.abs().max(); 
                    PTV_ST2 = gyro_moving2.abs().max();
                    Angle1=TurnAngle(gyro_moving1,Freq_LB)
                    Angle2=TurnAngle(gyro_moving2,Freq_LB)
                    turn_duration1=(idx_inversion-i_start_LB)/Freq_LB
                    turn_duration2=(i_end_LB-idx_inversion)/Freq_LB
                    
                    '''  
                    plt.figure()
                    plt.plot(t_LB, angle_vect, label='Accumulated Angle')
                    plt.axvline(t_LB[i_start_LB], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LB[idx_inversion], c= 'red', label='Inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Angle [deg]')
                    plt.title('Turn angle accumulated')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid()
                    ''' 
                    
                    # %% Feet part 
                    
                    gyro_ml_left=LF.gyr_z_LF
                    gyro_ml_right=RF.gyr_z_RF
                    
                    
                    th_LF = np.max(np.abs(gyro_ml_left)) * 0.2
                    th_RF = np.max(np.abs(gyro_ml_right)) * 0.2
                    
                    [duration_LF, i_start_LF, i_end_LF, gyro_moving_LF] = TurnDuration(gyro_ml_left,Freq_F, th_LF)
                    [duration_RF, i_start_RF, i_end_RF, gyro_moving_RF] = TurnDuration( gyro_ml_right,Freq_F, th_RF)
                    
                    if i_start_LF < i_start_RF: 
                        Initial_Foot = 'Left Foot'
                        i_start_F=i_start_LF
                    else:
                        Initial_Foot = 'Right Foot'
                        i_start_F=i_start_RF
                    
                    if i_end_LF > i_end_RF:
                        Final_Foot = 'Left Foot'
                        i_end_F=i_end_LF
                    else:
                        Final_Foot = 'Right Foot'
                        i_end_F = i_end_RF;
                        
                    duration_F=(i_end_F-i_start_F)/Freq_F
                    
                    [duration_LF1, i_start_LF1, i_end_LF1, gyro_moving_LF1] = TurnDuration( gyro_ml_left[:int(idx_inversion*Freq_F/Freq_LB)],Freq_F, th_LF)
                    [duration_RF1, i_start_RF1, i_end_RF1, gyro_moving_RF1] = TurnDuration( gyro_ml_right[:int(idx_inversion*Freq_F/Freq_LB)],Freq_F, th_RF)
                    [duration_LF2, i_start_LF2, i_end_LF2, gyro_moving_LF2] = TurnDuration( gyro_ml_left[int(idx_inversion*Freq_F/Freq_LB):],Freq_F, th_LF)
                    [duration_RF2, i_start_RF2, i_end_RF2, gyro_moving_RF2] = TurnDuration( gyro_ml_right[int(idx_inversion*Freq_F/Freq_LB):],Freq_F, th_RF)
                    
                    
                    if i_end_LF1 > i_end_RF1:
                        Final_Foot1 = 'Left Foot'
                        i_end_F1=i_end_LF1
                    else:
                        Final_Foot1 = 'Right Foot'
                        i_end_F1 = i_end_RF1;
                        
                    if i_start_LF2 < i_start_RF2: 
                        Initial_Foot2 = 'Left Foot'
                        i_start_F2=i_start_LF2
                    else:
                        Initial_Foot2 = 'Right Foot'
                        i_start_F2=i_start_RF2
                        
                    duration_intra_inversion=(i_start_F2-i_end_F1)/Freq_F
                    effective_duration=duration_F-duration_intra_inversion
                    
                    #%% number of steps
                
                    acc_z_left = LF.acc_z_LF    
                    
                    locsICs_LF, locsFCs_LF, gait_envelop_LF, gyro_f_LF=StepDetection( acc_z_left, gyro_ml_left, Freq_F, Orientation, idx_inversion)
                        
                    acc_z_right = RF.acc_z_RF
                     
                    locsICs_RF, locsFCs_RF, gait_envelop_RF, gyro_f_RF=StepDetection( acc_z_right, gyro_ml_right, Freq_F, Orientation, idx_inversion)
                    
                    scale = int(Freq_F / Freq_LB)
                    
                    mask_LF = (
                            ((locsICs_LF >= i_start_LB*scale-Freq_F) & (locsICs_LF <= i_end_LB*scale+Freq_F)) |
                            ((locsFCs_LF >= i_start_LB*scale-Freq_F) & (locsFCs_LF <= i_end_LB*scale+Freq_F))
                        )
                        
                    locsICs_LF = locsICs_LF[mask_LF]
                    locsFCs_LF = locsFCs_LF[mask_LF]
                    
                    mask_RF = (
                            ((locsICs_RF >= i_start_LB*scale-Freq_F) & (locsICs_RF <= i_end_LB*scale+Freq_F)) | 
                            ((locsFCs_RF >= i_start_LB*scale-Freq_F) & (locsFCs_RF <= i_end_LB*scale+Freq_F))
                        )
                        
                    locsICs_RF = locsICs_RF[mask_RF]
                    locsFCs_RF = locsFCs_RF[mask_RF]
                    
                    
                    numStepsLF=len(locsICs_LF)
                    numStepsRF=len(locsICs_RF)
                    
                    Tot_numSteps=numStepsLF + numStepsRF
                    
                    #print('The total number of steps:', Tot_numSteps, 'of which', numStepsLF, 'with the left foot and', numStepsRF, 'with the right foot\n')
                
                    numStepsLF1=sum(locsICs_LF<(idx_inversion*Freq_F/Freq_LB))
                    numStepsRF1=sum(locsICs_RF<(idx_inversion*Freq_F/Freq_LB))
                    numStepsLF2=sum(locsICs_LF>(idx_inversion*Freq_F/Freq_LB))
                    numStepsRF2=sum(locsICs_RF>(idx_inversion*Freq_F/Freq_LB))
                
                    
                    Tot_numSteps1=numStepsLF1 + numStepsRF1
                    Tot_numSteps2=numStepsLF2 + numStepsRF2
                    
                    
                    
                    # %% MTV, PTV and inversion
                    
                    MTV_LF = gyro_moving_LF.abs().mean()
                    PTV_LF = gyro_moving_LF.abs().max()
                    
                    MTV_RF = gyro_moving_RF.abs().mean()
                    PTV_RF = gyro_moving_RF.abs().max()
                    
                    MTV_FEET = np.mean([MTV_LF, MTV_RF])
                    PTV_FEET = np.max([PTV_LF, PTV_RF])
                    
                    
                    RMS_LF = np.sqrt(np.mean(gyro_moving_LF**2))
                    RMS_RF = np.sqrt(np.mean(gyro_moving_RF**2))
                    
                    if (RMS_RF + RMS_LF) != 0:
                        SI_RMS = (np.abs(RMS_RF - RMS_LF) / (0.5 * (RMS_RF + RMS_LF))) * 100
                    else:
                        SI_RMS = 0.0
                    
                    if (numStepsRF + numStepsLF) != 0:
                        SI_Steps = (np.abs(numStepsRF - numStepsLF) / (0.5 * (numStepsRF + numStepsLF))) * 100
                    else:
                        SI_Steps = 0.0
                    
                    idx_inversion_LF=int(idx_inversion*(Freq_F/Freq_LB))
                    idx_inversion_RF=idx_inversion_LF
                    
                    angle_RF=TurnAngle(gyro_moving_RF,Freq_F)
                    angle_LF=TurnAngle(gyro_moving_LF,Freq_F)
                    
                    gyro_moving_LF1 = gyro_moving_LF[:idx_inversion_LF-i_start_LF]; 
                    gyro_moving_LF2=gyro_moving_LF[idx_inversion_LF-i_start_LF+1:]; 
                    MTV_ST_LF1 = gyro_moving_LF1.abs().mean();
                    MTV_ST_LF2 = gyro_moving_LF2.abs().mean();
                    PTV_ST_LF1 = gyro_moving_LF1.abs().max(); 
                    PTV_ST_LF2 = gyro_moving_LF2.abs().max();
                    Angle_LF1=TurnAngle(gyro_moving_LF1,Freq_F)
                    Angle_LF2=TurnAngle(gyro_moving_LF2,Freq_F)
                    turn_duration_LF1=(idx_inversion_LF-i_start_LF)/Freq_F
                    turn_duration_LF2=(i_end_LF-idx_inversion_LF)/Freq_F
                    
                    gyro_moving_RF1 = gyro_moving_RF[:idx_inversion_RF-i_start_RF]; 
                    gyro_moving_RF2=gyro_moving_RF[idx_inversion_RF-i_start_RF+1:]; 
                    MTV_ST_RF1 = gyro_moving_RF1.abs().mean();
                    MTV_ST_RF2 = gyro_moving_RF2.abs().mean();
                    PTV_ST_RF1 = gyro_moving_RF1.abs().max(); 
                    PTV_ST_RF2 = gyro_moving_RF2.abs().max();
                    Angle_RF1=TurnAngle(gyro_moving_RF1,Freq_F)
                    Angle_RF2=TurnAngle(gyro_moving_RF2,Freq_F)
                    turn_duration_RF1=(idx_inversion_RF-i_start_RF)/Freq_F
                    turn_duration_RF2=(i_end_RF-idx_inversion_RF)/Freq_F
                
                    angle_F1 = np.mean([Angle_LF1, Angle_RF1])
                    angle_F2 = np.mean([Angle_LF2, Angle_RF2])
                    turn_duration1_F=(i_end_F1-i_start_F)/Freq_F
                    turn_duration2_F=(i_end_F-i_start_F2)/Freq_F
                    
                    # %% Table
                    
                    row = {
                        "Patient": code,
                        "Duration LB [s]": duration_LB,
                        "Duration turn 1 LB [s]": turn_duration1,
                        "Duration turn 2 LB [s]": turn_duration2,
                        "Duration F [s]": duration_F,
                        "Duration turn 1 F [s]": turn_duration1_F,
                        "Duration turn 2 F [s]": turn_duration2_F,
                        "Inversion Duration [s]": duration_intra_inversion,
                        "Effective Duration [s]": effective_duration,
                        "MTV [deg/s]": MTV,
                        "PTV [deg/s]": PTV,
                        "Total step": Tot_numSteps,
                        "Step with left foot": numStepsLF,
                        "Step with right foot": numStepsRF,
                        "Steps in the first turn": Tot_numSteps1,
                        "Steps in the second turn": Tot_numSteps2,
                        "RMS LF [deg/s]": RMS_LF,
                        "RMS RF [deg/s]": RMS_RF,
                        "Symmetry Index RMS [%]": SI_RMS,
                        "Symmetry Index Steps [%]": SI_Steps,
                        "Turn Angle 1 [deg]": Angle1,
                        "Turn Angle 2 [deg]": Angle2,
                        "Initial Foot": Initial_Foot,
                        "Direction 1st turn": Orientation
                    }
                    Table = pd.concat([Table, pd.DataFrame([row])], ignore_index=True)
                    # %% Plot 
                       
                    plt.figure()
                    plt.subplot(2, 1, 1)
                    plt.plot(t_LB, LB.acc_x, label='Ax')
                    plt.plot(t_LB, LB.acc_y, label='Ay')
                    plt.plot(t_LB, LB.acc_z, label='Az')
                    plt.axvline(t_LB[i_start_LB], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LB[idx_inversion], c= 'red', label='Inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Acceleration [m/sec^2]')
                    plt.title('Lower back linear acceleration')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    plt.subplot(2, 1, 2)
                    plt.plot(t_LB, LB.gyr_x, label='Gx')
                    plt.plot(t_LB, LB.gyr_y, label='Gy')
                    plt.plot(t_LB, LB.gyr_z, label='Gz')
                    plt.axvline(t_LB[i_start_LB], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LB[idx_inversion], c= 'red', label='Inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Angular velocity [deg/sec]')
                    plt.title('Lower back angular velocity')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    pdf.savefig() 
                    plt.close() 
                    
                    plt.figure()
                    plt.subplot(2, 1, 1)
                    plt.plot(t_LF, LF.acc_x_LF, label='Ax')
                    plt.plot(t_LF, LF.acc_y_LF, label='Ay')
                    plt.plot(t_LF, LF.acc_z_LF, label='Az')
                    plt.axvline(t_LF[i_start_LB*scale], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Acceleration [m/sec^2]')
                    plt.title('Left Foot linear acceleration')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    plt.subplot(2, 1, 2)
                    plt.plot(t_LF, LF.gyr_x_LF, label='Gx')
                    plt.plot(t_LF, LF.gyr_y_LF, label='Gy')
                    plt.plot(t_LF, LF.gyr_z_LF, label='Gz')
                    plt.axvline(t_LF[i_start_LB*scale],label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Angular velocity [deg/sec]')
                    plt.title('Left Foot angular velocity')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    pdf.savefig() 
                    plt.close() 
                    
                    
                    plt.figure()
                    plt.subplot(2, 1, 1)
                    plt.plot(t_RF, RF.acc_x_RF, label='Ax')
                    plt.plot(t_RF, RF.acc_y_RF, label='Ay')
                    plt.plot(t_RF, RF.acc_z_RF, label='Az')
                    plt.axvline(t_RF[i_start_LB*scale], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_RF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Acceleration [m/sec^2]')
                    plt.title('Right Foot linear acceleration')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    plt.subplot(2, 1, 2)
                    plt.plot(t_RF, RF.gyr_x_RF, label='Gx')
                    plt.plot(t_RF, RF.gyr_y_RF, label='Gy')
                    plt.plot(t_RF, RF.gyr_z_RF, label='Gz')
                    plt.axvline(t_RF[i_start_LB*scale], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_RF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Angular velocity [deg/sec]')
                    plt.title('Right Foot angular velocity')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    pdf.savefig() 
                    plt.close() 
                    
                    
                    plt.figure()
                    plt.plot(t_LB, angle_vect, label='Accumulated Angle')
                    plt.axvline(t_LB[i_start_LB], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_LB[idx_inversion], c= 'red', label='Inversion')
                    plt.xlabel('Time[s]')
                    plt.ylabel('Angle [deg]')
                    plt.title('Turn angle accumulated')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.legend(loc='upper right', fontsize=dim_legend)
                    plt.tight_layout()
                    plt.grid(True)
                    
                    pdf.savefig() 
                    plt.close() 
                    
                     
                    plt.figure()
                    plt.subplot(2, 1, 1)
                    plt.plot(t_LF,  gyro_f_LF)
                    for ic_t in t_LF[locsICs_LF]:
                        plt.axvline(ic_t, color='red', alpha=0.3)
                    for fc_t in t_LF[locsFCs_LF]:
                        plt.axvline(fc_t, color='green', alpha=0.3)
                    
                    plt.axvline(t_LB[idx_inversion], c= 'red')
                    plt.xlabel('Time [s]')
                    plt.title('Step detection LF')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.ylabel('Gyro_filt (LF)')
                    plt.axvline(t_RF[i_start_LB*scale], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_RF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.tight_layout()
                    plt.grid(True)
                     
                    plt.subplot(2, 1, 2)
                    plt.plot(t_LF,  gyro_f_RF)
                    for ic_t in t_LF[locsICs_RF]:
                        plt.axvline(ic_t, color='red', alpha=0.3)
                    for fc_t in t_LF[locsFCs_RF]:
                        plt.axvline(fc_t, color='green', alpha=0.3)
                    
                    plt.axvline(t_LB[idx_inversion], c= 'red')
                    plt.xlabel('Time [s]')
                    plt.title('Step detection RF')
                    plt.suptitle(f'Patient {code},{condition},{test},{trial}', fontsize=14)
                    plt.ylabel('Gyro_filt (RF)')
                    plt.axvline(t_RF[i_start_LB*scale], label='Start', c='red', linestyle='--')
                    plt.axvline(t_LB[i_end_LB], label='End', c='red', linestyle='-.')
                    plt.axvline(t_RF[idx_inversion*scale], c='red', label='Turn inversion')
                    plt.tight_layout()
                    plt.grid(True)
                    
                    pdf.savefig() 
                    plt.close() 
                
                # %% Save and modify the table                
                Table.to_excel(f"features_{condition}_{test}_{trial}.xlsx", index=False)
                
                
                file_path=f"features_{condition}_{test}_{trial}.xlsx"
                wb = load_workbook(file_path)
                ws = wb.active
                for col in ws.columns:
                    max_length = 0
                    col_letter = col[0].column_letter
                    
                    for cell in col:
                        try:
                            if cell.value:
                                max_length = max(max_length, len(str(cell.value)))
                        except:
                            pass
                    
                    ws.column_dimensions[col_letter].width = max_length + 2
                
                thin = Side(style="thin")
                border = Border(left=thin, right=thin, top=thin, bottom=thin)
                for row in ws.iter_rows(min_row=2):
                    for cell in row:
                        if isinstance(cell.value, float):
                            cell.number_format = "0.00"
                        cell.alignment = Alignment(horizontal="center")
                        cell.border = border
                
                
                wb.save(file_path)  
                
                
                
                
                
                
                
                
                