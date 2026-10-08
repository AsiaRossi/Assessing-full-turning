import pandas as pd
import numpy as np
import matplotlib.pyplot as plt
import seaborn as sns
from matplotlib.backends.backend_pdf import PdfPages
from openpyxl import load_workbook
from openpyxl.styles import Side, Border, Alignment, Font
from scipy.stats import ttest_rel, fisher_exact, spearmanr
from statsmodels.stats.multitest import multipletests

Features = [
    'Duration LB [s]',
    'Duration F [s]',
    'Duration turn 1 LB [s]',
    'Duration turn 1 F [s]',
    'Duration turn 2 LB [s]',
    'Duration turn 2 F [s]',
    'Inversion Duration [s]',
    'Effective Duration [s]',
    'MTV [deg/s]',
    'PTV [deg/s]',
    'Total step',
    'Steps in the first turn',
    'Steps in the second turn',
    'Step with left foot',
    'Step with right foot',
    'RMS LF [deg/s]',
    'RMS RF [deg/s]',
    'Symmetry Index RMS [%]',
    'Symmetry Index Steps [%]',
    'Turn Angle 1 [deg]',
    'Turn Angle 2 [deg]'
]

Files = {
    'Off_ST': {'t1': 'features_Condition2_Test5_Trial1.xlsx', 
               't2': 'features_Condition2_Test5_Trial2.xlsx'},
    'On_ST':  {'t1': 'features_Condition4_Test5_Trial1.xlsx',  
               't2': 'features_Condition4_Test5_Trial2.xlsx'},
    'Off_DT': {'t1': 'features_Condition2_Test6_Trial1.xlsx', 
               't2': 'features_Condition2_Test6_Trial2.xlsx'},
    'On_DT':  {'t1': 'features_Condition4_Test6_Trial1.xlsx',  
               't2': 'features_Condition4_Test6_Trial2.xlsx'},
}

PDF_Name = "Final plot reports.pdf"
Excel_Name = "Differences_and_Ttest_Results.xlsx"

def load_and_average(file_t1, file_t2, features, condition, task):
    try:
        df1_full = pd.read_excel(file_t1)
        df1 = df1_full[['Patient'] + features].copy()
        if 'Direction 1st turn' in df1_full.columns:
            df1['Direction 1st turn'] = df1_full['Direction 1st turn']
    except FileNotFoundError:
        print(f"Warning: File {file_t1} not found.")
        df1 = pd.DataFrame(columns=['Patient'] + features)

    try:
        df2_full = pd.read_excel(file_t2)
        df2 = df2_full[['Patient'] + features].copy()
    except FileNotFoundError:
        print(f"Warning: File {file_t2} not found.")
        df2 = pd.DataFrame(columns=['Patient'] + features)

    if df1.empty and df2.empty:
        return pd.DataFrame()

    merged = pd.merge(df1, df2, on='Patient', suffixes=('_t1', '_t2'), how='outer')
    df_avg = pd.DataFrame({'Patient': merged['Patient']})
    
    for feat in features:
        if f"{feat}_t1" in merged.columns and f"{feat}_t2" in merged.columns:
            df_avg[feat] = merged[[f"{feat}_t1", f"{feat}_t2"]].mean(axis=1, skipna=True)
        elif f"{feat}_t1" in merged.columns:
            df_avg[feat] = merged[f"{feat}_t1"]
        elif f"{feat}_t2" in merged.columns:
            df_avg[feat] = merged[f"{feat}_t2"]
        
    if 'Direction 1st turn' in merged.columns:
        df_avg['Direction 1st turn'] = merged['Direction 1st turn']
    
    df_avg['Medication'] = condition  
    df_avg['Task'] = task             
    
    return df_avg

def excel_formatting(file_path):
    try:
        wb = load_workbook(file_path)
        thin = Side(style="thin")
        border = Border(left=thin, right=thin, top=thin, bottom=thin)
        
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.border = border

            for col in ws.columns:
                max_length = 0
                col_letter = col[0].column_letter
                
                for cell in col:
                    try:
                        if cell.value is not None:
                            max_length = max(max_length, len(str(cell.value)))
                    except:
                        pass
                    
                    if cell.row > 1:
                        if isinstance(cell.value, float):
                            cell.number_format = "0.00"
                        cell.alignment = Alignment(horizontal="center")
                        cell.border = border
                
                ws.column_dimensions[col_letter].width = max_length + 2

        wb.save(file_path)
    except Exception as e:
        print(f"Error during excel formatting: {e}")

def run_paired_ttest(df, col1, col2, feature, condition, name1, name2):
    valid_data = df.dropna(subset=[col1, col2])
    if len(valid_data) > 1:
        mean1 = valid_data[col1].mean()
        sd1 = valid_data[col1].std()
        mean2 = valid_data[col2].mean()
        sd2 = valid_data[col2].std()
        
        t_stat, p_val = ttest_rel(valid_data[col1], valid_data[col2])
        
        return {
            'Feature': feature, 
            'Fixed Condition': condition, 
            'Group 1': name1,
            'Mean 1': mean1,
            'SD 1': sd1,
            'Group 2': name2,
            'Mean 2': mean2,
            'SD 2': sd2,
            't-stat': t_stat, 
            'p-value': p_val
        }
    return None

dfs = []
dfs.append(load_and_average(Files['Off_ST']['t1'], Files['Off_ST']['t2'], Features, 'Med Off', 'Single Task'))
dfs.append(load_and_average(Files['On_ST']['t1'], Files['On_ST']['t2'], Features, 'Med On', 'Single Task'))
dfs.append(load_and_average(Files['Off_DT']['t1'], Files['Off_DT']['t2'], Features, 'Med Off', 'Dual Task'))
dfs.append(load_and_average(Files['On_DT']['t1'], Files['On_DT']['t2'], Features, 'Med On', 'Dual Task'))

dfs = [df for df in dfs if not df.empty]
if len(dfs) < 4:
    print("Warning: Some groups are missing. PDF will be generated only with available data.")

final_data = pd.concat(dfs, ignore_index=True)
sns.set_theme(style="whitegrid", palette="muted")

off_st = next((df for df in dfs if df['Medication'].iloc[0] == 'Med Off' and df['Task'].iloc[0] == 'Single Task'), pd.DataFrame(columns=['Patient']+Features))
on_st  = next((df for df in dfs if df['Medication'].iloc[0] == 'Med On' and df['Task'].iloc[0] == 'Single Task'), pd.DataFrame(columns=['Patient']+Features))
off_dt = next((df for df in dfs if df['Medication'].iloc[0] == 'Med Off' and df['Task'].iloc[0] == 'Dual Task'), pd.DataFrame(columns=['Patient']+Features))
on_dt  = next((df for df in dfs if df['Medication'].iloc[0] == 'Med On' and df['Task'].iloc[0] == 'Dual Task'), pd.DataFrame(columns=['Patient']+Features))

#%% Med ON vs OFF
st_merge = pd.merge(off_st, on_st, on='Patient', suffixes=('_off', '_on'), how='outer')
dt_merge = pd.merge(off_dt, on_dt, on='Patient', suffixes=('_off', '_on'), how='outer')

excel_med = pd.DataFrame({'Patient': st_merge['Patient']}) if not st_merge.empty else pd.DataFrame(columns=['Patient'])
if not dt_merge.empty:
    excel_med = pd.merge(excel_med, pd.DataFrame({'Patient': dt_merge['Patient']}), on='Patient', how='outer')

ttest_results = []

for feat in Features:
    if f'{feat}_on' in st_merge.columns and f'{feat}_off' in st_merge.columns:
        excel_med[f'{feat} (Delta ST: On-Off)'] = st_merge[f'{feat}_on'] - st_merge[f'{feat}_off']
        res = run_paired_ttest(st_merge, f'{feat}_on', f'{feat}_off', feat, 'Single Task', 'Med On', 'Med Off')
        if res: ttest_results.append(res)
            
    if f'{feat}_on' in dt_merge.columns and f'{feat}_off' in dt_merge.columns:
        excel_med[f'{feat} (Delta DT: On-Off)'] = dt_merge[f'{feat}_on'] - dt_merge[f'{feat}_off']
        res = run_paired_ttest(dt_merge, f'{feat}_on', f'{feat}_off', feat, 'Dual Task', 'Med On', 'Med Off')
        if res: ttest_results.append(res)

#%% Single vs Dual Task
off_merge = pd.merge(off_st, off_dt, on='Patient', suffixes=('_st', '_dt'), how='outer')
on_merge = pd.merge(on_st, on_dt, on='Patient', suffixes=('_st', '_dt'), how='outer')

excel_task = pd.DataFrame({'Patient': off_merge['Patient']}) if not off_merge.empty else pd.DataFrame(columns=['Patient'])
if not on_merge.empty:
    excel_task = pd.merge(excel_task, pd.DataFrame({'Patient': on_merge['Patient']}), on='Patient', how='outer')

for feat in Features:
    if f'{feat}_dt' in off_merge.columns and f'{feat}_st' in off_merge.columns:
        excel_task[f'{feat} (Delta Off: DT-ST)'] = off_merge[f'{feat}_dt'] - off_merge[f'{feat}_st']
        res = run_paired_ttest(off_merge, f'{feat}_dt', f'{feat}_st', feat, 'Med Off', 'Dual Task', 'Single Task')
        if res: ttest_results.append(res)

    if f'{feat}_dt' in on_merge.columns and f'{feat}_st' in on_merge.columns:
        excel_task[f'{feat} (Delta On: DT-ST)'] = on_merge[f'{feat}_dt'] - on_merge[f'{feat}_st']
        res = run_paired_ttest(on_merge, f'{feat}_dt', f'{feat}_st', feat, 'Med On', 'Dual Task', 'Single Task')
        if res: ttest_results.append(res)

df_ttest = pd.DataFrame(ttest_results)

# Aggiunta p-value adjusted con Bonferroni
if not df_ttest.empty:
    _, p_adj, _, _ = multipletests(df_ttest['p-value'], method='bonferroni')
    df_ttest['p-value adjusted (Bonferroni)'] = p_adj

with pd.ExcelWriter(Excel_Name, engine='openpyxl') as writer:
    excel_med.to_excel(writer, sheet_name='Deltas_Med(On-Off)', index=False)
    excel_task.to_excel(writer, sheet_name='Deltas_Task(DT-ST)', index=False)
    if not df_ttest.empty:
        df_ttest.to_excel(writer, sheet_name='T_Test_Results', index=False)

excel_formatting(Excel_Name)

#%% Graphs 
with PdfPages(PDF_Name) as pdf:
    
    melted_data = pd.melt(
        final_data, 
        id_vars=['Patient', 'Medication', 'Task'], 
        value_vars=Features, 
        var_name='Feature', 
        value_name='Value'
    )

    # on vs off
    for feature in Features:
        data_feat = melted_data[melted_data['Feature'] == feature]
        if data_feat.empty: continue
        
        fig, ax = plt.subplots(figsize=(10, 6))
        sns.boxplot(
            data=data_feat, x='Task', y='Value', hue='Medication',   
            palette={'Med Off': '#fca311', 'Med On': '#00b4d8'}, width=0.6, fliersize=0, ax=ax
        )
        sns.swarmplot(
            data=data_feat, x='Task', y='Value', hue='Medication', dodge=True,          
            palette={'Med Off': '#000000', 'Med On': '#000000'}, alpha=0.6, size=4, legend=False, ax=ax
        )
        ax.set_title(f"{feature}: Med Off vs Med On in Single and Dual Task", fontsize=14, pad=15, fontweight='bold')
        ax.set_ylabel(f"Mean ({feature})", fontsize=12)
        ax.set_xlabel("Task", fontsize=12)
        ax.legend(loc='upper left', bbox_to_anchor=(1, 1))
        
        plt.tight_layout()
        pdf.savefig(fig, bbox_inches='tight')
        plt.close(fig)

    delta_long_med = pd.melt(excel_med, id_vars=['Patient'], var_name='Metric', value_name='Delta')
    delta_long_med['Task'] = delta_long_med['Metric'].apply(lambda x: 'Single Task' if 'ST' in x else 'Dual Task')
    delta_long_med['Feature'] = delta_long_med['Metric'].apply(lambda x: x.split(' (')[0])
    delta_long_med_clean = delta_long_med.dropna(subset=['Delta'])

    for feature in Features:
        data_feat = delta_long_med_clean[delta_long_med_clean['Feature'] == feature]
        if data_feat.empty: continue
            
        fig, ax = plt.subplots(figsize=(8, 6))
        sns.boxplot(data=data_feat, x='Task', y='Delta', width=0.5, color='#00b4d8', fliersize=0, ax=ax)
        sns.stripplot(data=data_feat, x='Task', y='Delta', color='black', alpha=0.7, jitter=True, size=6, ax=ax)
        ax.axhline(0, color='red', linestyle='--', linewidth=2, alpha=0.8)
        
        ax.set_title(f'Delta Plot (Med On - Med Off): {feature}', fontsize=14, fontweight='bold')
        ax.set_ylabel(f'Difference ({feature})', fontsize=12)
        ax.set_xlabel('Task', fontsize=12)
        
        plt.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)

    # ST vs DT
    for feature in Features:
        data_feat = melted_data[melted_data['Feature'] == feature]
        if data_feat.empty: continue
        
        fig, ax = plt.subplots(figsize=(10, 6))
        sns.boxplot(
            data=data_feat, x='Medication', y='Value', hue='Task',   
            palette={'Single Task': '#fca311', 'Dual Task': '#00b4d8'}, width=0.6, fliersize=0, ax=ax
        )
        sns.swarmplot(
            data=data_feat, x='Medication', y='Value', hue='Task', dodge=True,          
            palette={'Single Task': '#000000', 'Dual Task': '#000000'}, alpha=0.6, size=4, legend=False, ax=ax
        )
        ax.set_title(f"{feature}: Single Task vs Dual Task in Med Off and Med On", fontsize=14, pad=15, fontweight='bold')
        ax.set_ylabel(f"Mean ({feature})", fontsize=12)
        ax.set_xlabel("Medication", fontsize=12)
        ax.legend(loc='upper left', bbox_to_anchor=(1, 1))
        
        plt.tight_layout()
        pdf.savefig(fig, bbox_inches='tight')
        plt.close(fig)

    delta_long_task = pd.melt(excel_task, id_vars=['Patient'], var_name='Metric', value_name='Delta')
    delta_long_task['Medication'] = delta_long_task['Metric'].apply(lambda x: 'Med Off' if 'Off' in x else 'Med On')
    delta_long_task['Feature'] = delta_long_task['Metric'].apply(lambda x: x.split(' (')[0])
    delta_long_task_clean = delta_long_task.dropna(subset=['Delta'])

    for feature in Features:
        data_feat = delta_long_task_clean[delta_long_task_clean['Feature'] == feature]
        if data_feat.empty: continue
            
        fig, ax = plt.subplots(figsize=(8, 6))
        sns.boxplot(data=data_feat, x='Medication', y='Delta', width=0.5, color='#00b4d8', fliersize=0, ax=ax)
        sns.stripplot(data=data_feat, x='Medication', y='Delta', color='black', alpha=0.7, jitter=True, size=6, ax=ax)
        ax.axhline(0, color='red', linestyle='--', linewidth=2, alpha=0.8)
        
        ax.set_title(f'Delta Plot (Dual Task - Single Task): {feature}', fontsize=14, fontweight='bold')
        ax.set_ylabel(f'Difference ({feature})', fontsize=12)
        ax.set_xlabel('Medication', fontsize=12)
        
        plt.tight_layout()
        pdf.savefig(fig)
        plt.close(fig)
        
        
#%% Correlation direction 1st turn - side of onset

file_paths = {
    'Off_ST_t1': Files['Off_ST']['t1'],
    'Off_ST_t2': Files['Off_ST']['t2'],
    'On_ST_t1': Files['On_ST']['t1'],
    'On_ST_t2': Files['On_ST']['t2'],
    'Off_DT_t1': Files['Off_DT']['t1'],
    'Off_DT_t2': Files['Off_DT']['t2'],
    'On_DT_t1': Files['On_DT']['t1'],
    'On_DT_t2': Files['On_DT']['t2']
}

df_directions = None

for col_name, path in file_paths.items():
    try:
        temp_df = pd.read_excel(path)
        if 'Patient' in temp_df.columns and 'Direction 1st turn' in temp_df.columns:
            temp_df = temp_df[['Patient', 'Direction 1st turn']].copy()
            temp_df.rename(columns={'Direction 1st turn': col_name}, inplace=True)
            
            if df_directions is None:
                df_directions = temp_df
            else:
                df_directions = pd.merge(df_directions, temp_df, on='Patient', how='outer')
    except FileNotFoundError:
        pass 

if df_directions is not None:
    dir_columns = [col for col in df_directions.columns if col != 'Patient']
    
    def check_maintenance(row):
        vals = row[dir_columns].dropna().unique()
        if len(vals) == 1:
            return 'Keeped'
        elif len(vals) > 1:
            return 'Changed'
        return np.nan
        
    df_directions['Comportamento Direzione'] = df_directions.apply(check_maintenance, axis=1)
    
    def get_mode(row):
        vals = row[dir_columns].dropna()
        if len(vals) > 0:
            return vals.mode().iloc[0] 
        return np.nan
        
    df_directions['Direzione Prevalente'] = df_directions.apply(get_mode, axis=1)

    try:
        df_lista = pd.read_excel('ListaID.xlsx')
        
        id_col = next((col for col in df_lista.columns if 'id' in col.lower() or 'patient' in col.lower()), None)
        if id_col:
            df_lista.rename(columns={id_col: 'Patient'}, inplace=True)
            
        if 'LATO PEGGIORE' in df_lista.columns:
            df_final_dir = pd.merge(df_directions, df_lista[['Patient', 'LATO PEGGIORE']], on='Patient', how='left')
            
            mapping_lato = {'sx': 'Left', 'dx': 'Right'}
            df_final_dir['Lato Peggiore Mappato'] = df_final_dir['LATO PEGGIORE'].str.strip().str.lower().map(mapping_lato)
            
            contingency_table = pd.crosstab(df_final_dir['Direzione Prevalente'], df_final_dir['Lato Peggiore Mappato'])
            
            print("\n--- Tabella di Contingenza (Direzione vs Lato Peggiore) ---")
            print(contingency_table)
            
            if contingency_table.shape == (2, 2):
                oddsratio, p_value = fisher_exact(contingency_table)
                print(f"\nTest di Fisher - p-value: {p_value:.4f}")
                if p_value < 0.05:
                    print("Result: There is a statistically significant correlation.")
                else:
                    print("Result: There is NO statistically significant correlation (p >= 0.05).")
            else:
                print("\nWarning: There is not enough data in either category (Left/Right) to perform Fisher's exact test on the 2x2 matrix.")
                
            Excel_Dir_Name = "Tracking_Direction_Correlation.xlsx"
            with pd.ExcelWriter(Excel_Dir_Name, engine='openpyxl') as writer:
                df_final_dir.to_excel(writer, sheet_name='Analisi_Direzioni', index=False)
                contingency_table.to_excel(writer, sheet_name='Contingence')
                
            excel_formatting(Excel_Dir_Name)
            print(f"\nFinito! Tutti i dati sulle direzioni salvati in: {Excel_Dir_Name}")
            
        else:
            print("Colonna 'LATO PEGGIORE' non trovata in ListaID.xlsx.")
            
    except FileNotFoundError:
        print("Errore: Il file 'ListaID.xlsx' non è stato trovato nella directory.")
else:
    print("Nessun dato sulle direzioni trovato negli 8 file.")
    
#%% Correlation features - UPDRS 

try:
    df_lista = pd.read_excel('ListaID.xlsx')
    
    id_col = next((col for col in df_lista.columns if 'id' in col.lower() or 'patient' in col.lower()), None)
    if id_col:
        df_lista.rename(columns={id_col: 'Patient'}, inplace=True)

    df_lista['UPDRS OFF'] = pd.to_numeric(df_lista['UPDRS OFF'], errors='coerce')
    df_lista['UPDRS ON'] = pd.to_numeric(df_lista['UPDRS ON'], errors='coerce')

    df_lista_clean = df_lista.dropna(subset=['UPDRS OFF', 'UPDRS ON'])

    df_corr = pd.merge(final_data, df_lista_clean[['Patient', 'UPDRS OFF', 'UPDRS ON']], on='Patient', how='inner')

    def match_updrs(row):
        if row['Medication'] == 'Med Off':
            return row['UPDRS OFF']
        elif row['Medication'] == 'Med On':
            return row['UPDRS ON']
        return np.nan

    df_corr['MDS_UPDRSIII'] = df_corr.apply(match_updrs, axis=1)
    df_corr = df_corr.dropna(subset=['MDS_UPDRSIII'])

    tasks = ['Single Task', 'Dual Task']
    
    PDF_Heatmap_Name = "Correlazioni_UPDRS_Heatmap.pdf"
    
    with PdfPages(PDF_Heatmap_Name) as pdf:
        for task in tasks:
            df_task = df_corr[df_corr['Task'] == task]
            if df_task.empty:
                continue
            
            cols_to_correlate = Features + ['MDS_UPDRSIII']
            df_numeric = df_task[cols_to_correlate].apply(pd.to_numeric, errors='coerce')
            
            corr_matrix = df_numeric.corr(method='spearman')
            
            updrs_corr = corr_matrix[['MDS_UPDRSIII']].drop('MDS_UPDRSIII')
            updrs_corr = updrs_corr.sort_values(by='MDS_UPDRSIII', ascending=False)
            
            fig, ax = plt.subplots(figsize=(8, 10))
            sns.heatmap(updrs_corr, 
                        annot=True,       
                        fmt=".2f",        
                        cmap="coolwarm",  
                        vmin=-1, vmax=1,  
                        cbar_kws={'label': 'Correlation Coefficient (Spearman)'},
                        linewidths=0.5,
                        ax=ax)
            
            ax.set_title(f"Correlation between Features and UPDRS\n({task})", fontsize=14, fontweight='bold', pad=15)
            ax.set_ylabel("Cinematic Features")
            
            plt.tight_layout()
            pdf.savefig(fig)
            plt.close(fig)
            
            fig2, ax2 = plt.subplots(figsize=(16, 12))
            sns.heatmap(corr_matrix, annot=False, cmap="coolwarm", vmin=-1, vmax=1, ax=ax2)
            ax2.set_title(f"Complete Correalation Matrix - {task}", fontsize=16, fontweight='bold', pad=15)
            plt.tight_layout()
            pdf.savefig(fig2)
            plt.close(fig2)

    print(f"Grafici Heatmap generati e salvati in: {PDF_Heatmap_Name}")

except FileNotFoundError:
    print("Errore: Il file 'ListaID.xlsx' non è stato trovato.")
except Exception as e:
    print(f"Si è verificato un errore durante l'analisi delle correlazioni: {e}")  

if 'df_corr' in locals():
    risultati_correlazione = []
    tasks = ['Single Task', 'Dual Task']
    
    matrici_correlazione = {}
    
    for task in tasks:
        df_task = df_corr[df_corr['Task'] == task]
        if df_task.empty: continue
        
        for feat in Features:
            valid_data = df_task.dropna(subset=[feat, 'MDS_UPDRSIII'])
            
            if len(valid_data) > 2: 
                rho, p_val = spearmanr(valid_data[feat], valid_data['MDS_UPDRSIII'])
                
                risultati_correlazione.append({
                    'Task': task,
                    'Feature': feat,
                    'Spearman Rho': rho,
                    'p-value': p_val
                })
                
        df_numeric_features = df_task[Features].apply(pd.to_numeric, errors='coerce')
        matrice_corr = df_numeric_features.corr(method='spearman')
        matrici_correlazione[task] = matrice_corr
                
    df_risultati_corr = pd.DataFrame(risultati_correlazione)
    
    if not df_risultati_corr.empty:
        # Calcolo p-value corretto con Bonferroni per ciascun task
        df_risultati_corr['p-value adjusted (Bonferroni)'] = np.nan
        for task in tasks:
            idx = df_risultati_corr['Task'] == task
            p_vals = df_risultati_corr.loc[idx, 'p-value']
            if len(p_vals) > 0:
                _, p_adj, _, _ = multipletests(p_vals, method='bonferroni')
                df_risultati_corr.loc[idx, 'p-value adjusted (Bonferroni)'] = p_adj

        df_risultati_corr['Significativo (p_adj<0.05)'] = df_risultati_corr['p-value adjusted (Bonferroni)'].apply(lambda x: 'Si' if x < 0.05 else 'No')

        Excel_Corr_Name = "Tabella_Correlazioni_Completa.xlsx" 
        
        with pd.ExcelWriter(Excel_Corr_Name, engine='openpyxl') as writer:
            df_st = df_risultati_corr[df_risultati_corr['Task'] == 'Single Task'].drop(columns=['Task'])
            df_dt = df_risultati_corr[df_risultati_corr['Task'] == 'Dual Task'].drop(columns=['Task'])
            
            df_st = df_st.sort_values(by='p-value')
            df_dt = df_dt.sort_values(by='p-value')
            
            df_st.to_excel(writer, sheet_name='UPDRS_ST', index=False)
            df_dt.to_excel(writer, sheet_name='UPDRS_DT', index=False)
            
            if 'Single Task' in matrici_correlazione:
                df_matrice_st = matrici_correlazione['Single Task'].reset_index().rename(columns={'index': 'Feature'})
                df_matrice_st.to_excel(writer, sheet_name='Features_vs_Features_ST', index=False)
            
            if 'Dual Task' in matrici_correlazione:
                df_matrice_dt = matrici_correlazione['Dual Task'].reset_index().rename(columns={'index': 'Feature'})
                df_matrice_dt.to_excel(writer, sheet_name='Features_vs_Features_DT', index=False)
            
        excel_formatting(Excel_Corr_Name)
        print(f"Tabella salvata con successo in: {Excel_Corr_Name}")
    else:
        print("Non ci sono dati sufficienti per calcolare le correlazioni in tabella.")
else:
    print("Errore: il dataframe df_corr non è stato trovato. Assicurati di eseguire prima il blocco della heatmap.")
    
#%% Analisi Durata Turn: Lato Migliore vs Lato Peggiore

try:
    df_lista = pd.read_excel('ListaID.xlsx')
    id_col = next((col for col in df_lista.columns if 'id' in col.lower() or 'patient' in col.lower()), None)
    if id_col:
        df_lista.rename(columns={id_col: 'Patient'}, inplace=True)

    def norm_dir(val):
        if pd.isna(val): 
            return None
        s = str(val).strip().lower()
        if s in ['sx', 'left', 'sinistra']: 
            return 'Left'
        if s in ['dx', 'right', 'destra']: 
            return 'Right'
        return None

    df_lista['Worse_Side'] = df_lista['LATO PEGGIORE'].apply(norm_dir)
    df_turn_sides = df_lista[['Patient']].drop_duplicates().copy()

    for cond_key, trials in Files.items():
        for trial_key, file_path in trials.items():
            col_prefix = f"{cond_key}_{trial_key}"
            
            try:
                df_raw = pd.read_excel(file_path)
                
                col_t1 = 'Duration turn 1 LB [s]' if 'Duration turn 1 LB [s]' in df_raw.columns else ('Duration turn 1 F [s]' if 'Duration turn 1 F [s]' in df_raw.columns else None)
                col_t2 = 'Duration turn 2 LB [s]' if 'Duration turn 2 LB [s]' in df_raw.columns else ('Duration turn 2 F [s]' if 'Duration turn 2 F [s]' in df_raw.columns else None)
                
                if 'Patient' in df_raw.columns and 'Direction 1st turn' in df_raw.columns and col_t1 and col_t2:
                    
                    df_sub = pd.merge(
                        df_raw[['Patient', 'Direction 1st turn', col_t1, col_t2]], 
                        df_lista[['Patient', 'Worse_Side']], 
                        on='Patient', 
                        how='inner'
                    )
                    
                    dur_better_list, dur_worse_list = [], []
                    
                    for idx, row in df_sub.iterrows():
                        dir_t1 = norm_dir(row['Direction 1st turn'])
                        worse = row['Worse_Side']
                        
                        if dir_t1 and worse:
                            better = 'Right' if worse == 'Left' else 'Left'
                            dur1, dur2 = row[col_t1], row[col_t2]
                            
                            if dir_t1 == better:
                                dur_better, dur_worse = dur1, dur2
                            elif dir_t1 == worse:
                                dur_better, dur_worse = dur2, dur1
                            else:
                                dur_better, dur_worse = np.nan, np.nan
                        else:
                            dur_better, dur_worse = np.nan, np.nan
                            
                        dur_better_list.append(dur_better)
                        dur_worse_list.append(dur_worse)
                    
                    col_b = f"{col_prefix} - Turn Lato Migliore [s]"
                    col_w = f"{col_prefix} - Turn Lato Peggiore [s]"
                    
                    df_sub[col_b] = dur_better_list
                    df_sub[col_w] = dur_worse_list
                    
                    df_turn_sides = pd.merge(df_turn_sides, df_sub[['Patient', col_b, col_w]], on='Patient', how='left')
                    
            except Exception as e:
                print(f"Errore caricamento {file_path}: {e}")

    df_turn_avg = pd.DataFrame({'Patient': df_turn_sides['Patient']})
    conditions = ['Off_ST', 'On_ST', 'Off_DT', 'On_DT']
    turn_features = ['Turn Lato Migliore [s]', 'Turn Lato Peggiore [s]']

    for cond in conditions:
        for feat in turn_features:
            col_t1 = f"{cond}_t1 - {feat}"
            col_t2 = f"{cond}_t2 - {feat}"
            if col_t1 in df_turn_sides.columns and col_t2 in df_turn_sides.columns:
                df_turn_avg[f"{cond} - {feat}"] = df_turn_sides[[col_t1, col_t2]].mean(axis=1, skipna=True)

    def compute_ttest(df, col1, col2, feature, condition_label, g1_label, g2_label):
        valid_data = df.dropna(subset=[col1, col2])
        if len(valid_data) > 1:
            m1, s1 = valid_data[col1].mean(), valid_data[col1].std()
            m2, s2 = valid_data[col2].mean(), valid_data[col2].std()
            t_stat, p_val = ttest_rel(valid_data[col1], valid_data[col2])
            return {
                'Feature': feature,
                'Condizione Fissata': condition_label,
                'Gruppo 1': g1_label,
                'Media 1 [s]': m1,
                'Dev.Std 1 [s]': s1,
                'Gruppo 2': g2_label,
                'Media 2 [s]': m2,
                'Dev.Std 2 [s]': s2,
                'Statistica t': t_stat,
                'p-value': p_val
            }
        return None

    ttest_results = []

    for feat in turn_features:
        # In Med OFF
        res = compute_ttest(df_turn_avg, f"Off_DT - {feat}", f"Off_ST - {feat}", feat, 'Med OFF', 'Dual Task', 'Single Task')
        if res: ttest_results.append(res)
        # In Med ON
        res = compute_ttest(df_turn_avg, f"On_DT - {feat}", f"On_ST - {feat}", feat, 'Med ON', 'Dual Task', 'Single Task')
        if res: ttest_results.append(res)

    for feat in turn_features:
        # In Single Task
        res = compute_ttest(df_turn_avg, f"On_ST - {feat}", f"Off_ST - {feat}", feat, 'Single Task', 'Med ON', 'Med OFF')
        if res: ttest_results.append(res)
        # In Dual Task
        res = compute_ttest(df_turn_avg, f"On_DT - {feat}", f"Off_DT - {feat}", feat, 'Dual Task', 'Med ON', 'Med OFF')
        if res: ttest_results.append(res)

    df_ttest_summary = pd.DataFrame(ttest_results)
    if not df_ttest_summary.empty:
        _, p_adj, _, _ = multipletests(df_ttest_summary['p-value'], method='bonferroni')
        df_ttest_summary['p-value adjusted (Bonferroni)'] = p_adj

    ttest_lato_results = []

    for cond in conditions:
        col_best = f"{cond} - Turn Lato Migliore [s]"
        col_worse = f"{cond} - Turn Lato Peggiore [s]"
        
        valid_data = df_turn_avg.dropna(subset=[col_best, col_worse])
        if len(valid_data) > 1:
            m_best, s_best = valid_data[col_best].mean(), valid_data[col_best].std()
            m_worse, s_worse = valid_data[col_worse].mean(), valid_data[col_worse].std()
            t_stat, p_val = ttest_rel(valid_data[col_best], valid_data[col_worse])
            
            ttest_lato_results.append({
                'Condizione': cond,
                'Media Lato Migliore [s]': m_best,
                'Dev.Std Lato Migliore [s]': s_best,
                'Media Lato Peggiore [s]': m_worse,
                'Dev.Std Lato Peggiore [s]': s_worse,
                'Statistica t': t_stat,
                'p-value': p_val
            })

    df_ttest_lato = pd.DataFrame(ttest_lato_results)
    if not df_ttest_lato.empty:
        _, p_adj, _, _ = multipletests(df_ttest_lato['p-value'], method='bonferroni')
        df_ttest_lato['p-value adjusted (Bonferroni)'] = p_adj

    Excel_Turn_Sides_Name = "Durata_bestvsworst.xlsx"
    with pd.ExcelWriter(Excel_Turn_Sides_Name, engine='openpyxl') as writer:
        df_turn_sides.to_excel(writer, sheet_name='Dati_Per_Trial', index=False)
        df_turn_avg.to_excel(writer, sheet_name='Medie_Pazienti', index=False)
        df_ttest_summary.to_excel(writer, sheet_name='Statistiche_Confronti', index=False)
        df_ttest_lato.to_excel(writer, sheet_name='Confronto', index=False)

    excel_formatting(Excel_Turn_Sides_Name)
except Exception as e:
    print(f"Errore durante l'analisi: {e}")
